from types import SimpleNamespace

from openpyxl import load_workbook

from app.models.cms import TourCategory
from app.services import tour_import_export as tour_import


class _Query:
    def __init__(self, *, rows=None, first=None):
        self._rows = rows or []
        self._first = first

    def filter(self, *_args):
        return self

    def order_by(self, *_args):
        return self

    def all(self):
        return self._rows

    def first(self):
        return self._first


class _TemplateDb:
    def query(self, model):
        assert model is TourCategory
        return _Query()


def test_supplier_template_uses_operating_currency_and_explains_lock():
    workbook = load_workbook(
        tour_import.build_import_template_workbook(
            _TemplateDb(), include_supplier_column=False, supplier_currency="nzd"
        )
    )

    tours = workbook["Tours"]
    headers = [cell.value for cell in tours[1]]
    currency_index = headers.index("Currency") + 1
    assert "Supplier (admin only)" not in headers
    assert tours.cell(row=2, column=currency_index).value == "NZD"
    assert "NZD" in [row[2] for row in workbook["Reference"].iter_rows(min_row=2, values_only=True)]
    assert "Currency is locked to your operating currency (NZD)" in workbook["Instructions"]["A3"].value


def test_supplier_import_uses_profile_currency_not_spreadsheet_value(monkeypatch):
    saved = []

    monkeypatch.setattr(tour_import, "_supplier_operating_currency", lambda _db, _supplier_id: "NZD")
    monkeypatch.setattr(
        tour_import,
        "save_tour",
        lambda _db, payload, _actor, _request: saved.append(payload) or {"id": 77, "title": payload.title},
    )

    result = tour_import.import_tours(
        db=SimpleNamespace(),
        rows=[{"title": "Imported NZ tour", "currency": "USD", "days": 3}],
        actor=SimpleNamespace(),
        actor_supplier_id=18,
    )

    assert result == [{"row": 2, "status": "created", "tour_id": 77, "title": "Imported NZ tour"}]
    assert saved[0].supplier_id == 18
    assert saved[0].currency == "NZD"


def test_generated_template_can_be_parsed_back_into_import_rows():
    template = tour_import.build_import_template_workbook(_TemplateDb(), include_supplier_column=False, supplier_currency="NZD")
    rows = tour_import.parse_tour_import_rows(template.getvalue())

    assert rows == [{
        "title": "Golden Triangle Explorer",
        "subtitle": "Delhi - Agra - Jaipur in 6 days",
        "category": "Cultural",
        "country": "India",
        "state": "Delhi",
        "city": "New Delhi",
        "start_location": "New Delhi",
        "finish_location": "Jaipur",
        "currency": "NZD",
        "days": 6,
        "hours": None,
        "nights": 5,
        "max_group_size": 15,
        "min_booking_size": 2,
        "tour_language": "English",
        "suitable_age_range": "All ages",
        "visibility": "public",
        "short_description": "A classic introduction to India's cultural heartland.",
    }]

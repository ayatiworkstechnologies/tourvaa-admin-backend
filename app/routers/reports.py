import csv
import io
import json
import logging
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session, joinedload

from app.database import get_db
from app.models.agents import Agent
from app.models.audit import AuditLog
from app.services.audit import log_audit
from app.models.bookings import Booking
from app.models.cms import City, Country, Tour, TourCategory
from app.models.cancellations import CancellationRequest
from app.models.reviews import TourReview
from app.models.supplier_ledger import SupplierLedger, SupplierPayout, SupplierPayoutItem
from app.models.agent_ledger import AgentLedger
from app.auth.permissions import require_any_permission
from app.utils.money import money_str, utcnow
from app.models.customers import Customer
from app.models.invoices import Invoice
from app.models.payments import Payment
from app.models.reports import ReportSchedule
from app.models.suppliers import Supplier
from app.models.users import User
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

router = APIRouter(prefix="/reports", tags=["Reports"])
logger = logging.getLogger(__name__)

PERIODS = {"day", "week", "month", "quarter", "half_year", "year", "custom", "all"}


def _scheduled_report_count(db: Session) -> int:
    # report_schedules (migration 20260728_0050) may not exist yet on every
    # deployed database; degrade to 0 instead of a 500 until that migration
    # has run everywhere.
    try:
        return db.query(func.count(ReportSchedule.id)).filter(ReportSchedule.is_active == True).scalar() or 0  # noqa: E712
    except SQLAlchemyError:
        logger.exception("Failed to count report schedules")
        return 0


def _report_role(actor: User | None) -> str:
    if not actor or not actor.role:
        return "admin"
    slug = actor.role.slug or ""
    if "supplier" in slug:
        return "supplier"
    if "agent" in slug:
        return "agent"
    return "admin"


def _scope_bookings(query, db: Session, actor: User | None):
    """Restrict a Booking-based query to the caller's own tenant for supplier/agent roles."""
    role = _report_role(actor)
    if role == "supplier":
        supplier = db.query(Supplier).filter(Supplier.user_id == actor.id).first()
        return query.filter(Booking.supplier_id == (supplier.id if supplier else -1))
    if role == "agent":
        agent = db.query(Agent).filter(Agent.user_id == actor.id).first()
        return query.filter(Booking.agent_id == (agent.id if agent else -1))
    return query


def _require_admin_report(actor: User | None):
    if _report_role(actor) != "admin":
        raise HTTPException(status_code=403, detail="This report is restricted to administrators")


def _money(value):
    return money_str(value or 0)


NON_SALES_BOOKING_STATUSES = ("cancelled", "declined", "refunded")


def _sales_reconciliation(gross, settled, outstanding) -> dict[str, Decimal]:
    """Return one booking-period sales reconciliation.

    A financial month must be reconciled from a single set of bookings:
    gross booking value = settled amount + outstanding amount.  Invoice rows
    are documents (and may be generated per payment), so their total is kept
    separate and must never be added to outstanding balances.
    """
    gross_value = Decimal(str(gross or 0))
    settled_value = Decimal(str(settled or 0))
    outstanding_value = Decimal(str(outstanding or 0))
    return {
        "gross_sales": gross_value,
        "settled_amount": settled_value,
        "outstanding_amount": outstanding_value,
        "reconciliation_difference": gross_value - settled_value - outstanding_value,
    }


def _supplier_booking_value(b: Booking):
    """What a booking is worth at the supplier's own prices (the locked
    snapshot's supplier gross), for supplier-facing reports -- the customer
    total includes Tourvaa's admin markup, which suppliers never see.
    Bookings without a snapshot carried no markup, so final_amount is used."""
    if b.supplier_gross_amount is not None:
        return b.supplier_gross_amount
    return b.final_amount


def _booking_amount_column(actor: User | None):
    """Per-booking amount to sum in a report: supplier gross for supplier
    actors (see _supplier_booking_value), customer final amount otherwise."""
    if _report_role(actor) == "supplier":
        return func.coalesce(Booking.supplier_gross_amount, Booking.final_amount)
    return Booking.final_amount


def _amount_to_usd(amount, currency: str | None) -> float:
    """Convert mixed booking currencies to the reporting base before summing."""
    if not amount:
        return 0.0
    from app.services.currency import BASE_CURRENCY, convert_amount, normalize_currency
    normalized = normalize_currency(currency, BASE_CURRENCY)
    if normalized == BASE_CURRENCY:
        return float(amount)
    converted, _, _ = convert_amount(Decimal(str(amount)), normalized, BASE_CURRENCY)
    return float(converted)


def _period_range(period: str, start_date: str = "", end_date: str = ""):
    """Return (start, end) datetimes for the given calendar-aligned period, or
    None for a bound that shouldn't be filtered (e.g. "all", or a custom bound
    left blank)."""
    now = utcnow()
    period = (period or "all").strip().lower()

    if period == "custom":
        start = _parse_date(start_date) if start_date else None
        end = _parse_date(end_date, end_of_day=True) if end_date else None
        return start, end

    if period == "day":
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    elif period == "week":
        today = now.replace(hour=0, minute=0, second=0, microsecond=0)
        start = today - timedelta(days=today.weekday())
    elif period == "month":
        start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    elif period == "quarter":
        quarter_start_month = ((now.month - 1) // 3) * 3 + 1
        start = now.replace(month=quarter_start_month, day=1, hour=0, minute=0, second=0, microsecond=0)
    elif period == "half_year":
        half_start_month = 1 if now.month <= 6 else 7
        start = now.replace(month=half_start_month, day=1, hour=0, minute=0, second=0, microsecond=0)
    elif period == "year":
        start = now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
    else:
        return None, None

    return start, now


def _parse_date(value: str, end_of_day: bool = False):
    dt = datetime.strptime(value.strip(), "%Y-%m-%d").replace(tzinfo=timezone.utc)
    if end_of_day:
        dt = dt.replace(hour=23, minute=59, second=59)
    return dt


def _apply_range(query, column, start, end):
    if start is not None:
        query = query.filter(column >= start)
    if end is not None:
        query = query.filter(column <= end)
    return query


def _period_params(
    period: str = Query(default="all", description="day|week|month|quarter|half_year|year|custom|all"),
    start_date: str = Query(default=""),
    end_date: str = Query(default=""),
):
    if period not in PERIODS:
        raise HTTPException(status_code=400, detail=f"period must be one of {sorted(PERIODS)}")
    start, end = _period_range(period, start_date, end_date)
    return {"period": period, "start": start, "end": end}


@router.get("/summary")
def summary(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view"))):
    start, end = params["start"], params["end"]
    bookings_q = _scope_bookings(_apply_range(db.query(Booking), Booking.created_at, start, end), db, current_user)
    sales_bookings_q = bookings_q.filter(Booking.booking_status.notin_(NON_SALES_BOOKING_STATUSES))

    total_bookings = bookings_q.with_entities(func.count(Booking.id)).scalar() or 0
    confirmed = bookings_q.filter(Booking.booking_status == "confirmed").with_entities(func.count(Booking.id)).scalar() or 0
    cancelled = bookings_q.filter(Booking.booking_status == "cancelled").with_entities(func.count(Booking.id)).scalar() or 0
    reconciliation = _sales_reconciliation(
        sales_bookings_q.with_entities(func.coalesce(func.sum(Booking.final_amount), 0)).scalar(),
        sales_bookings_q.with_entities(func.coalesce(func.sum(Booking.amount_paid), 0)).scalar(),
        sales_bookings_q.with_entities(func.coalesce(func.sum(Booking.amount_pending), 0)).scalar(),
    )
    # Scope invoice documents through the same booking period. An invoice may
    # represent a deposit or balance installment, so this is audit metadata,
    # not a component of gross sales and is intentionally not reconciled.
    invoice_q = _scope_bookings(
        _apply_range(db.query(Invoice).join(Booking, Invoice.booking_id == Booking.id), Booking.created_at, start, end),
        db,
        current_user,
    )
    invoice_total = invoice_q.with_entities(func.coalesce(func.sum(Invoice.total_amount), 0)).scalar() or 0
    return {
        "status": "success",
        "data": {
            "total_bookings": total_bookings,
            "confirmed_bookings": confirmed,
            "cancelled_bookings": cancelled,
            # Retained for existing API consumers; these values now share the
            # gross-sales booking scope instead of payment-created-at scope.
            "captured_revenue": _money(reconciliation["settled_amount"]),
            "pending_payments": _money(reconciliation["outstanding_amount"]),
            "gross_sales": _money(reconciliation["gross_sales"]),
            "settled_amount": _money(reconciliation["settled_amount"]),
            "outstanding_amount": _money(reconciliation["outstanding_amount"]),
            "reconciliation_difference": _money(reconciliation["reconciliation_difference"]),
            "invoice_total": _money(invoice_total),
        },
    }


@router.get("/bookings")
def booking_report(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    query = _scope_bookings(_apply_range(db.query(Booking), Booking.created_at, params["start"], params["end"]), db, current_user)
    rows = query.with_entities(Booking.booking_status, Booking.currency, _booking_amount_column(current_user)).all()
    agg: dict[str, dict] = {}
    for status, currency, amount in rows:
        bucket = agg.setdefault(status, {"count": 0, "amount": 0.0})
        bucket["count"] += 1
        bucket["amount"] += _amount_to_usd(amount, currency)
    return {"status": "success", "data": [{"status": status, "count": v["count"], "amount": _money(v["amount"])} for status, v in agg.items()]}


@router.get("/payments")
def payment_report(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    query = _apply_range(db.query(Payment), Payment.created_at, params["start"], params["end"])
    if _report_role(current_user) != "admin":
        query = _scope_bookings(query.join(Booking, Payment.booking_id == Booking.id), db, current_user)
    rows = query.with_entities(Payment.payment_status, func.count(Payment.id), func.coalesce(func.sum(Payment.captured_amount), 0), func.coalesce(func.sum(Payment.refunded_amount), 0)).group_by(Payment.payment_status).all()
    return {"status": "success", "data": [{"status": status, "count": count, "captured": _money(captured), "refunded": _money(refunded)} for status, count, captured, refunded in rows]}


@router.get("/pending-payments")
def pending_payments(db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    query = _scope_bookings(db.query(Booking).filter(Booking.amount_pending > 0), db, current_user)
    rows = query.order_by(Booking.amount_pending.desc()).limit(200).all()
    return {"status": "success", "data": [{"booking_id": b.id, "booking_code": b.booking_code, "customer_id": b.customer_id, "amount_pending": _money(b.amount_pending), "payment_status": b.payment_status} for b in rows]}


@router.get("/overdue-payments")
def overdue_payments(db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    today = utcnow().date()
    query = _scope_bookings(db.query(Booking).filter(Booking.amount_pending > 0, Booking.tour_start_date != None, func.date(Booking.tour_start_date) <= today), db, current_user)
    rows = query.order_by(Booking.tour_start_date.asc()).limit(200).all()
    return {"status": "success", "data": [{"booking_id": b.id, "booking_code": b.booking_code, "tour_start_date": b.tour_start_date, "amount_pending": _money(b.amount_pending)} for b in rows]}


@router.get("/country-wise")
def country_wise(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    query = db.query(Country.country_name, Booking.id, Booking.currency, Booking.final_amount).join(Booking, Booking.country_id == Country.id, isouter=True)
    query = _apply_range(query, Booking.created_at, params["start"], params["end"])
    query = _scope_bookings(query, db, current_user)
    agg: dict[str, dict] = {}
    for country, booking_id, currency, amount in query.all():
        bucket = agg.setdefault(country, {"count": 0, "amount": 0.0})
        if booking_id is not None:
            bucket["count"] += 1
            bucket["amount"] += _amount_to_usd(amount, currency)
    return {"status": "success", "data": [{"country": country, "bookings": v["count"], "amount": _money(v["amount"])} for country, v in agg.items()]}


@router.get("/cancellations")
def cancellations(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    query = db.query(Booking).filter(Booking.booking_status == "cancelled")
    query = _apply_range(query, Booking.cancelled_at, params["start"], params["end"])
    query = _scope_bookings(query, db, current_user)
    rows = query.order_by(Booking.cancelled_at.desc()).limit(200).all()
    return {"status": "success", "data": [{"booking_id": b.id, "booking_code": b.booking_code, "reason": b.cancellation_reason, "cancelled_at": b.cancelled_at, "amount": _money(b.final_amount)} for b in rows]}


@router.get("/suppliers")
def supplier_report(
    params: dict = Depends(_period_params),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.supplier", "reports.admin")),
):
    _require_admin_report(current_user)
    query = db.query(Supplier.id, Supplier.supplier_name, Booking.id, Booking.currency, Booking.final_amount).join(Booking, Booking.supplier_id == Supplier.id, isouter=True)
    query = _apply_range(query, Booking.created_at, params["start"], params["end"])
    agg: dict[int, dict] = {}
    for sid, name, booking_id, currency, amount in query.all():
        bucket = agg.setdefault(sid, {"name": name, "count": 0, "amount": 0.0})
        if booking_id is not None:
            bucket["count"] += 1
            bucket["amount"] += _amount_to_usd(amount, currency)
    return {"status": "success", "data": [{"supplier_id": sid, "supplier_name": v["name"], "bookings": v["count"], "amount": _money(v["amount"])} for sid, v in agg.items()]}


@router.get("/agents")
def agent_report(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.agent", "reports.admin"))):
    _require_admin_report(current_user)
    query = db.query(Agent.id, Agent.agent_name, Booking.id, Booking.currency, Booking.final_amount).join(Booking, Booking.agent_id == Agent.id, isouter=True)
    query = _apply_range(query, Booking.created_at, params["start"], params["end"])
    agg: dict[int, dict] = {}
    for aid, name, booking_id, currency, amount in query.all():
        bucket = agg.setdefault(aid, {
            "name": name,
            "count": 0,
            "amount": 0.0,
            "commission_earned": 0.0,
            "commission_paid": 0.0,
            "commission_outstanding": 0.0,
        })
        if booking_id is not None:
            bucket["count"] += 1
            bucket["amount"] += _amount_to_usd(amount, currency)

    # The ledger is the source of truth for the commission promised to, paid
    # to, and still owed to an agent.  Keep its period aligned to the booking
    # report so the booking and commission figures always describe the same
    # set of bookings.  A reversed ledger row is retained for audit history,
    # but no longer represents earned/outstanding commission.
    ledger_query = db.query(
        AgentLedger.agent_id,
        AgentLedger.currency,
        AgentLedger.commission_amount,
        AgentLedger.amount_paid,
        AgentLedger.amount_pending,
        AgentLedger.status,
    ).join(Booking, AgentLedger.booking_id == Booking.id)
    ledger_query = _apply_range(ledger_query, Booking.created_at, params["start"], params["end"])
    for aid, currency, commission, paid, outstanding, status in ledger_query.all():
        bucket = agg.setdefault(aid, {
            "name": None,
            "count": 0,
            "amount": 0.0,
            "commission_earned": 0.0,
            "commission_paid": 0.0,
            "commission_outstanding": 0.0,
        })
        if status != "refunded":
            bucket["commission_earned"] += _amount_to_usd(commission, currency)
        bucket["commission_paid"] += _amount_to_usd(paid, currency)
        bucket["commission_outstanding"] += _amount_to_usd(outstanding, currency)

    return {
        "status": "success",
        "data": [
            {
                "agent_id": aid,
                "agent_name": v["name"] or f"Agent #{aid}",
                "bookings": v["count"],
                "amount": _money(v["amount"]),
                "commission_earned": _money(v["commission_earned"]),
                "commission_paid": _money(v["commission_paid"]),
                "commission_outstanding": _money(v["commission_outstanding"]),
            }
            for aid, v in agg.items()
        ],
    }


@router.get("/customers")
def customer_report(params: dict = Depends(_period_params), db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view", "reports.admin"))):
    _require_admin_report(current_user)
    query = db.query(Customer.id, Customer.full_name, Booking.id, Booking.currency, Booking.final_amount, Booking.amount_pending).join(Booking, Booking.customer_id == Customer.id, isouter=True)
    query = _apply_range(query, Booking.created_at, params["start"], params["end"])
    agg: dict[int, dict] = {}
    for cid, name, booking_id, currency, amount, pending in query.all():
        bucket = agg.setdefault(cid, {"name": name, "count": 0, "amount": 0.0, "pending": 0.0})
        if booking_id is not None:
            bucket["count"] += 1
            bucket["amount"] += _amount_to_usd(amount, currency)
            bucket["pending"] += _amount_to_usd(pending, currency)
    return {"status": "success", "data": [{"customer_id": cid, "customer_name": v["name"], "bookings": v["count"], "amount": _money(v["amount"]), "pending": _money(v["pending"])} for cid, v in agg.items()]}


def _audience_rows(db: Session, audience: str, entity_id: int | None, params: dict, actor: User) -> list[dict]:
    if audience == "supplier":
        rows = supplier_report(params=params, db=db, current_user=actor)["data"]
        id_key = "supplier_id"
    elif audience == "agent":
        rows = agent_report(params=params, db=db, current_user=actor)["data"]
        id_key = "agent_id"
    elif audience == "customer":
        rows = customer_report(params=params, db=db, current_user=actor)["data"]
        id_key = "customer_id"
    else:
        raise HTTPException(status_code=400, detail="audience must be supplier, agent, or customer")
    rows = [row for row in rows if entity_id is None or row[id_key] == entity_id]
    ids = [row[id_key] for row in rows]
    if not ids:
        return rows

    if audience == "supplier":
        entities = db.query(Supplier).options(joinedload(Supplier.user)).filter(Supplier.id.in_(ids)).all()
        details = {
            row.id: {
                "code": row.supplier_code,
                "email": row.user.email if row.user else None,
                "phone": (row.user.mobile_number or row.user.phone) if row.user else None,
                "status": row.status,
                "approval_status": row.approval_status,
                "joined_at": row.created_at,
            }
            for row in entities
        }
    elif audience == "agent":
        entities = db.query(Agent).options(joinedload(Agent.user)).filter(Agent.id.in_(ids)).all()
        details = {
            row.id: {
                "code": row.agent_code,
                "email": row.user.email if row.user else None,
                "phone": (row.user.mobile_number or row.user.phone) if row.user else None,
                "status": row.status,
                "approval_status": row.approval_status,
                "joined_at": row.created_at,
            }
            for row in entities
        }
    else:
        entities = db.query(Customer).filter(Customer.id.in_(ids)).all()
        details = {
            row.id: {
                "code": row.customer_code,
                "email": row.email,
                "phone": f"{row.phone_country_code} {row.phone}".strip(),
                "status": "blocked" if row.is_blocked else row.status,
                "approval_status": "verified" if row.email_verified else "unverified",
                "joined_at": row.created_at,
            }
            for row in entities
        }

    for row in rows:
        row.update(details.get(row[id_key], {}))
    return rows


@router.get("/audience-options")
def audience_options(
    audience: str = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    if audience == "supplier":
        rows = db.query(Supplier.id, Supplier.supplier_name).order_by(Supplier.supplier_name.asc()).all()
    elif audience == "agent":
        rows = db.query(Agent.id, Agent.agent_name).order_by(Agent.agent_name.asc()).all()
    elif audience == "customer":
        rows = db.query(Customer.id, Customer.full_name).order_by(Customer.full_name.asc()).all()
    else:
        raise HTTPException(status_code=400, detail="audience must be supplier, agent, or customer")
    return {"status": "success", "data": [{"id": row[0], "name": row[1] or f"{audience.title()} #{row[0]}"} for row in rows]}


@router.get("/audience")
def audience_report(
    audience: str = Query(...),
    entity_id: int | None = Query(default=None, ge=1),
    params: dict = Depends(_period_params),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    if params["start"] and params["end"] and params["start"] > params["end"]:
        raise HTTPException(status_code=400, detail="Start date must be on or before end date")
    rows = _audience_rows(db, audience, entity_id, params, current_user)
    return {"status": "success", "data": rows, "meta": {"audience": audience, "entity_id": entity_id, "scope": "selected" if entity_id else "all", "count": len(rows)}}


@router.get("/audience-detail")
def audience_report_detail(
    audience: str = Query(...),
    entity_id: int = Query(..., ge=1),
    params: dict = Depends(_period_params),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    """Return the booking-level evidence behind one audience report row."""
    _require_admin_report(current_user)
    if audience not in {"supplier", "agent", "customer"}:
        raise HTTPException(status_code=400, detail="audience must be supplier, agent, or customer")
    if params["start"] and params["end"] and params["start"] > params["end"]:
        raise HTTPException(status_code=400, detail="Start date must be on or before end date")

    id_column = {
        "supplier": Booking.supplier_id,
        "agent": Booking.agent_id,
        "customer": Booking.customer_id,
    }[audience]
    query = (
        db.query(Booking)
        .options(joinedload(Booking.customer))
        .filter(id_column == entity_id)
        .order_by(Booking.created_at.desc(), Booking.id.desc())
    )
    query = _apply_range(query, Booking.created_at, params["start"], params["end"])
    bookings = query.all()

    identity_rows = _audience_rows(db, audience, entity_id, params, current_user)
    identity = identity_rows[0] if identity_rows else {f"{audience}_id": entity_id, f"{audience}_name": f"{audience.title()} #{entity_id}"}
    totals = {
        "bookings": len(bookings),
        "travellers": 0,
        "booking_value": Decimal("0"),
        "paid": Decimal("0"),
        "pending": Decimal("0"),
        "partner_earnings": Decimal("0"),
        "tourvaa_revenue": Decimal("0"),
    }
    status_breakdown: dict[str, int] = {}
    payment_breakdown: dict[str, int] = {}
    detail_rows = []

    for booking in bookings:
        booking_value = _amount_to_usd(booking.final_amount, booking.currency)
        paid = _amount_to_usd(booking.amount_paid, booking.currency)
        pending = _amount_to_usd(booking.amount_pending, booking.currency)
        if audience == "supplier":
            partner_earnings = _amount_to_usd(booking.supplier_net_payable, booking.currency)
        elif audience == "agent":
            partner_earnings = _amount_to_usd(booking.agent_commission_amount, booking.currency)
        else:
            partner_earnings = paid
        tourvaa_revenue = _amount_to_usd(
            booking.tourvaa_net_revenue if booking.tourvaa_net_revenue is not None else booking.tourvaa_gross_revenue,
            booking.currency,
        )
        travellers = int(booking.total_travellers or (booking.no_of_adults or 0) + (booking.no_of_children or 0) + (booking.no_of_infants or 0))

        totals["travellers"] += travellers
        totals["booking_value"] += Decimal(str(booking_value))
        totals["paid"] += Decimal(str(paid))
        totals["pending"] += Decimal(str(pending))
        totals["partner_earnings"] += Decimal(str(partner_earnings))
        totals["tourvaa_revenue"] += Decimal(str(tourvaa_revenue))
        status_breakdown[booking.booking_status] = status_breakdown.get(booking.booking_status, 0) + 1
        payment_breakdown[booking.payment_status] = payment_breakdown.get(booking.payment_status, 0) + 1
        detail_rows.append({
            "booking_id": booking.id,
            "booking_code": booking.booking_code or f"Booking #{booking.id}",
            "booked_at": booking.created_at,
            "tour_name": booking.tour_name,
            "travel_date": booking.tour_start_date or booking.tour_date,
            "customer_name": booking.customer.full_name if booking.customer else None,
            "travellers": travellers,
            "booking_status": booking.booking_status,
            "payment_status": booking.payment_status,
            "currency": booking.currency,
            "booking_value": money_str(booking.final_amount),
            "paid": money_str(booking.amount_paid),
            "pending": money_str(booking.amount_pending),
            "partner_earnings": money_str(booking.supplier_net_payable if audience == "supplier" else booking.agent_commission_amount if audience == "agent" else booking.amount_paid),
            "tourvaa_revenue": money_str(booking.tourvaa_net_revenue if booking.tourvaa_net_revenue is not None else booking.tourvaa_gross_revenue),
        })

    return {
        "status": "success",
        "data": {
            "identity": identity,
            "summary": {key: (value if isinstance(value, int) else money_str(value)) for key, value in totals.items()},
            "status_breakdown": [{"status": key, "count": value} for key, value in sorted(status_breakdown.items())],
            "payment_breakdown": [{"status": key, "count": value} for key, value in sorted(payment_breakdown.items())],
            "bookings": detail_rows,
        },
    }


@router.get("/audience-export")
def audience_report_export(
    audience: str = Query(...),
    entity_id: int | None = Query(default=None, ge=1),
    period: str = Query(default="custom"),
    start_date: str = Query(default=""),
    end_date: str = Query(default=""),
    format: str = Query(default="xlsx", pattern="^(xlsx|pdf|csv)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.export", "reports.admin")),
):
    _require_admin_report(current_user)
    params = _period_params(period, start_date, end_date)
    rows = _audience_rows(db, audience, entity_id, params, current_user)
    columns = [
        (f"{audience.title()} ID", f"{audience}_id"), ("Code", "code"),
        ("Name", f"{audience}_name"), ("Email", "email"), ("Phone", "phone"),
        ("Status", "status"), ("Approval", "approval_status"), ("Joined", "joined_at"),
        ("Bookings", "bookings"), ("Booking Value (USD)", "amount"),
    ]
    if audience == "agent":
        columns += [("Commission Earned (USD)", "commission_earned"), ("Commission Paid (USD)", "commission_paid"), ("Commission Outstanding (USD)", "commission_outstanding")]
    elif audience == "customer":
        columns += [("Pending (USD)", "pending")]

    total_bookings = sum(int(row.get("bookings") or 0) for row in rows)
    total_value = sum(Decimal(str(row.get("amount") or 0)) for row in rows)
    scope = rows[0].get(f"{audience}_name") if entity_id and rows else f"All {audience.title()}s"
    title = f"Tourvaa {audience.title()} Report"
    metadata = [
        ("Scope", scope or f"{audience.title()} #{entity_id}"),
        ("Period", f"{start_date or 'All time'} to {end_date or 'Present'}"),
        ("Records", len(rows)), ("Total bookings", total_bookings),
        ("Total booking value", f"USD {total_value:,.2f}"),
        ("Generated", utcnow().strftime("%d %b %Y, %H:%M UTC")),
    ]
    basename = f"{audience}-report-{start_date or 'all'}-{end_date or 'all'}"

    if format == "csv":
        text_buffer = io.StringIO()
        writer = csv.writer(text_buffer)
        writer.writerow([title])
        for label, value in metadata:
            writer.writerow([label, value])
        writer.writerow([])
        writer.writerow([label for label, _ in columns])
        writer.writerows([[row.get(key, "") for _, key in columns] for row in rows])
        return StreamingResponse(iter([text_buffer.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{basename}.csv"'})

    binary_buffer = io.BytesIO()
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = f"{audience.title()} Report"
        last_column = get_column_letter(len(columns))
        sheet.merge_cells(f"A1:{last_column}1")
        sheet["A1"] = title
        sheet["A1"].font = Font(size=20, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="175CD3")
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34
        for index, (label, value) in enumerate(metadata, start=3):
            sheet.cell(index, 1, label).font = Font(bold=True, color="475467")
            sheet.cell(index, 2, value)
        header_row = 10
        for column_index, (label, _) in enumerate(columns, start=1):
            cell = sheet.cell(header_row, column_index, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="344054")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        thin = Side(style="thin", color="E4E7EC")
        for row_index, row in enumerate(rows, start=header_row + 1):
            for column_index, (_, key) in enumerate(columns, start=1):
                value = row.get(key, "")
                if isinstance(value, datetime):
                    value = value.strftime("%d %b %Y")
                cell = sheet.cell(row_index, column_index, value)
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
        widths = [14, 16, 25, 30, 18, 14, 16, 16, 12] + [20] * (len(columns) - 9)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(header_row, header_row + len(rows))}"
        sheet.sheet_view.showGridLines = False
        workbook.save(binary_buffer)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        extension = "xlsx"
    else:
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("report-title", parent=styles["Title"], textColor=colors.HexColor("#175CD3"), fontSize=20, leading=24, spaceAfter=8)
        meta_style = ParagraphStyle("report-meta", parent=styles["Normal"], textColor=colors.HexColor("#475467"), fontSize=8, leading=11)
        cell_style = ParagraphStyle("report-cell", parent=styles["Normal"], fontSize=6.5, leading=8, textColor=colors.HexColor("#101828"))
        document = SimpleDocTemplate(binary_buffer, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
        story = [Paragraph(title, title_style)]
        story.append(Table(
            [[Paragraph(f"<b>{escape(str(label))}</b><br/>{escape(str(value))}", meta_style) for label, value in metadata]],
            colWidths=[44 * mm] * 6,
            style=TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F4F7")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D5DD")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]),
        ))
        story.append(Spacer(1, 6 * mm))
        table_data = [[Paragraph(f"<b>{escape(label)}</b>", cell_style) for label, _ in columns]]
        for row in rows:
            table_data.append([Paragraph(escape(str(row.get(key, "") or "-")), cell_style) for _, key in columns])
        table = Table(table_data, colWidths=[277 * mm / len(columns)] * len(columns), repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#344054")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D0D5DD")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(table)
        document.build(story)
        media_type = "application/pdf"
        extension = "pdf"

    binary_buffer.seek(0)
    return StreamingResponse(iter([binary_buffer.getvalue()]), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{basename}.{extension}"'})


ROW_REPORT_LIMIT = 500


def _adult_child_counts(b: Booking) -> tuple[int, int]:
    # Two parallel pairs of columns exist on Booking - adults_count/children_count
    # is the one services/bookings.py prefers when both are populated.
    adults = b.adults_count if b.adults_count is not None else b.no_of_adults
    children = b.children_count if b.children_count is not None else b.no_of_children
    return adults or 0, children or 0


@router.get("/booking-report")
def booking_detail_report(
    params: dict = Depends(_period_params),
    booking_status: str = Query(default=""),
    payment_status: str = Query(default=""),
    country_id: int | None = Query(default=None),
    supplier_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    query = _apply_range(
        db.query(Booking).options(joinedload(Booking.customer), joinedload(Booking.agent), joinedload(Booking.tour), joinedload(Booking.supplier)),
        Booking.created_at, params["start"], params["end"],
    )
    if booking_status:
        query = query.filter(Booking.booking_status == booking_status.strip().lower())
    if payment_status:
        query = query.filter(Booking.payment_status == payment_status.strip().lower())
    if country_id:
        query = query.filter(Booking.country_id == country_id)
    if supplier_id:
        query = query.filter(Booking.supplier_id == supplier_id)
    rows = query.order_by(Booking.created_at.desc()).limit(ROW_REPORT_LIMIT).all()

    data = []
    for b in rows:
        adults, children = _adult_child_counts(b)
        data.append({
            "booking_id": b.id,
            "booking_code": b.booking_code,
            "booking_date": b.created_at,
            "customer_or_agent": (b.agent.agent_name if b.agent else None) or (b.customer.full_name if b.customer else None) or "-",
            "tour_name": b.tour.title if b.tour else b.tour_name,
            "supplier": b.supplier.supplier_name if b.supplier else b.supplier_name,
            "travel_date": b.tour_start_date or b.tour_date,
            "adults": adults,
            "children": children,
            "booking_amount": _money(b.final_amount),
            "payment_status": b.payment_status,
            "booking_status": b.booking_status,
            "cancellation_status": "cancelled" if b.booking_status == "cancelled" else "-",
        })
    return {"status": "success", "data": data}


@router.get("/sales-revenue-report")
def sales_revenue_report(
    params: dict = Depends(_period_params),
    granularity: str = Query(default="day", description="day|week|month"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    start, end = params["start"], params["end"]
    booking_q = _apply_range(db.query(Booking), Booking.created_at, start, end)
    payment_q = _apply_range(db.query(Payment).filter(Payment.payment_status.notin_(["voided", "failed"])), Payment.created_at, start, end)

    total_bookings = booking_q.with_entities(func.count(Booking.id)).scalar() or 0
    gross_booking_value = float(booking_q.with_entities(func.coalesce(func.sum(Booking.final_amount), 0)).scalar() or 0)
    discounts = float(booking_q.with_entities(func.coalesce(func.sum(Booking.discount_amount), 0)).scalar() or 0)
    taxes = float(booking_q.with_entities(func.coalesce(func.sum(Booking.tax_amount), 0)).scalar() or 0)

    # Ledger entries can be booked in a supplier's own currency (Supplier.currency
    # is never forced to USD, unlike Tour/TourPricing), so a platform-wide total
    # must convert each row to USD before summing rather than blending raw
    # amounts across currencies -- see services.supplier_ledger.get_supplier_statement
    # for the same concern on a per-supplier statement.
    from app.services.currency import BASE_CURRENCY, convert_amount, normalize_currency

    ledger_q = db.query(SupplierLedger).join(Booking, SupplierLedger.booking_id == Booking.id)
    ledger_q = _apply_range(ledger_q, Booking.created_at, start, end)
    platform_commission = 0.0
    supplier_payable = 0.0
    for row in ledger_q.with_entities(SupplierLedger.currency, SupplierLedger.commission_amount, SupplierLedger.net_payable):
        row_currency = normalize_currency(row.currency, BASE_CURRENCY)
        commission_amount = row.commission_amount or 0
        net_payable = row.net_payable or 0
        if row_currency != BASE_CURRENCY:
            if commission_amount:
                commission_amount, _, _ = convert_amount(Decimal(str(commission_amount)), row_currency, BASE_CURRENCY)
            if net_payable:
                net_payable, _, _ = convert_amount(Decimal(str(net_payable)), row_currency, BASE_CURRENCY)
        platform_commission += float(commission_amount)
        supplier_payable += float(net_payable)

    captured = float(payment_q.with_entities(func.coalesce(func.sum(Payment.captured_amount), 0)).scalar() or 0)
    refunded = float(payment_q.with_entities(func.coalesce(func.sum(Payment.refunded_amount), 0)).scalar() or 0)
    net_platform_revenue = max(0.0, captured - refunded)

    # Daily/weekly/monthly sales breakdown, bucketed on the payment date.
    if granularity not in {"day", "week", "month"}:
        raise HTTPException(status_code=400, detail="granularity must be one of day, week, month")
    if granularity == "day":
        bucket = func.date(Payment.created_at)
    elif granularity == "week":
        bucket = func.date_format(Payment.created_at, "%Y-%u")
    else:
        bucket = func.date_format(Payment.created_at, "%Y-%m")
    series_rows = (
        payment_q.with_entities(bucket.label("bucket"), func.coalesce(func.sum(Payment.captured_amount), 0))
        .group_by(bucket)
        .order_by(bucket)
        .all()
    )
    time_series = [{"period": str(bucket_value), "sales": _money(amount)} for bucket_value, amount in series_rows]

    return {
        "status": "success",
        "data": {
            "total_bookings": total_bookings,
            "gross_booking_value": _money(gross_booking_value),
            "discounts": _money(discounts),
            "taxes": _money(taxes),
            "platform_commission": _money(platform_commission),
            "supplier_payable": _money(supplier_payable),
            "refund_amount": _money(refunded),
            "net_platform_revenue": _money(net_platform_revenue),
            "granularity": granularity,
            "time_series": time_series,
        },
    }


@router.get("/payment-report")
def payment_detail_report(
    params: dict = Depends(_period_params),
    payment_status: str = Query(default=""),
    gateway: str = Query(default=""),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    query = _apply_range(db.query(Payment), Payment.created_at, params["start"], params["end"])
    if payment_status:
        query = query.filter(Payment.payment_status == payment_status.strip().lower())
    if gateway:
        query = query.filter(Payment.gateway == gateway.strip().lower())
    rows = query.order_by(Payment.created_at.desc()).limit(ROW_REPORT_LIMIT).all()

    data = []
    for p in rows:
        refund_status = "refunded" if p.payment_status == "refunded" else "partially_refunded" if p.payment_status == "partially_refunded" else "none"
        data.append({
            "transaction_id": p.transaction_id or p.gateway_payment_id,
            "booking_id": p.booking_id,
            "payment_gateway": p.gateway,
            "payment_method": p.payment_method,
            "paid_amount": _money(p.captured_amount or p.paid_amount),
            "payment_date": p.payment_date or p.created_at,
            "payment_status": p.payment_status,
            "failed_payment_reason": p.failure_reason,
            "refund_status": refund_status,
        })
    return {"status": "success", "data": data}


@router.get("/supplier-report")
def supplier_detail_report(
    params: dict = Depends(_period_params),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    suppliers = db.query(Supplier).all()
    data = []
    for s in suppliers:
        booking_q = _apply_range(db.query(Booking).filter(Booking.supplier_id == s.id), Booking.created_at, params["start"], params["end"])
        total_bookings = booking_q.with_entities(func.count(Booking.id)).scalar() or 0
        gross_sales = booking_q.with_entities(func.coalesce(func.sum(Booking.final_amount), 0)).scalar() or 0
        active_tours = db.query(func.count(Tour.id)).filter(Tour.supplier_id == s.id, Tour.status == "published").scalar() or 0

        ledger_q = _apply_range(db.query(SupplierLedger).filter(SupplierLedger.supplier_id == s.id), SupplierLedger.created_at, params["start"], params["end"])
        commission_deducted = ledger_q.with_entities(func.coalesce(func.sum(SupplierLedger.commission_amount), 0)).scalar() or 0
        amount_paid = ledger_q.with_entities(func.coalesce(func.sum(SupplierLedger.amount_paid), 0)).scalar() or 0
        outstanding_payable = ledger_q.with_entities(func.coalesce(func.sum(SupplierLedger.amount_pending), 0)).scalar() or 0

        data.append({
            "supplier_name": s.supplier_name,
            # No distinct "company name" field exists on Supplier - the
            # supplier's own name IS the company name in this data model.
            "company_name": s.supplier_name,
            "registration_date": s.created_at,
            "verification_status": s.approval_status,
            "active_tours": active_tours,
            "total_bookings": total_bookings,
            "gross_sales": _money(gross_sales),
            "commission_deducted": _money(commission_deducted),
            "amount_paid": _money(amount_paid),
            "outstanding_payable": _money(outstanding_payable),
            "supplier_status": s.status,
        })
    return {"status": "success", "data": data}


@router.get("/supplier-payout-report")
def supplier_payout_report(
    status: str = Query(default=""),
    supplier_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    query = db.query(SupplierPayout).options(joinedload(SupplierPayout.supplier))
    if status:
        query = query.filter(SupplierPayout.status == status.strip().lower())
    if supplier_id:
        query = query.filter(SupplierPayout.supplier_id == supplier_id)
    payouts = query.order_by(SupplierPayout.created_at.desc()).limit(ROW_REPORT_LIMIT).all()

    data = []
    for payout in payouts:
        items = (
            db.query(SupplierLedger)
            .join(SupplierPayoutItem, SupplierPayoutItem.ledger_id == SupplierLedger.id)
            .filter(SupplierPayoutItem.payout_id == payout.id)
            .options(joinedload(SupplierLedger.booking))
            .all()
        )
        total_booking_amount = sum(float(item.gross_amount or 0) for item in items)
        commission = sum(float(item.commission_amount or 0) for item in items)
        booking_dates = [item.booking.created_at for item in items if item.booking and item.booking.created_at]
        payout_period = f"{min(booking_dates).date()} - {max(booking_dates).date()}" if booking_dates else str(payout.created_at.date())

        data.append({
            "payout_id": payout.payout_code or payout.id,
            "supplier": payout.supplier.supplier_name if payout.supplier else None,
            "payout_period": payout_period,
            "total_booking_amount": _money(total_booking_amount),
            "commission": _money(commission),
            # No adjustment/refund-deduction line-item model exists on
            # supplier payouts today - a refund against an already-paid
            # ledger row is only flagged in a free-text note (see
            # services/supplier_ledger.py reverse_ledger_entry), not
            # itemized as a numeric deduction, so these are honestly 0
            # rather than a fabricated figure.
            "refund_deductions": _money(0),
            "other_adjustments": _money(0),
            "net_payable": _money(payout.total_amount),
            "payment_reference": payout.reference_number,
            "payout_status": payout.status,
            "paid_date": payout.paid_at,
        })
    return {"status": "success", "data": data}


@router.get("/tour-performance-report")
def tour_performance_report(
    supplier_id: int | None = Query(default=None),
    category_id: int | None = Query(default=None),
    country_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    query = db.query(Tour).options(joinedload(Tour.supplier), joinedload(Tour.country), joinedload(Tour.city), joinedload(Tour.category))
    if supplier_id:
        query = query.filter(Tour.supplier_id == supplier_id)
    if category_id:
        query = query.filter(Tour.category_id == category_id)
    if country_id:
        query = query.filter(Tour.country_id == country_id)
    tours = query.order_by(Tour.id.desc()).limit(ROW_REPORT_LIMIT).all()

    data = []
    for tour in tours:
        booking_q = db.query(Booking).filter(Booking.tour_id == tour.id)
        total_bookings = booking_q.with_entities(func.count(Booking.id)).scalar() or 0
        cancelled_bookings = booking_q.filter(Booking.booking_status == "cancelled").with_entities(func.count(Booking.id)).scalar() or 0
        confirmed_travellers = (
            booking_q.filter(Booking.booking_status.in_(["confirmed", "ongoing", "completed"]))
            .with_entities(func.coalesce(func.sum(Booking.total_travellers), 0))
            .scalar() or 0
        )
        revenue = booking_q.with_entities(func.coalesce(func.sum(Booking.final_amount), 0)).scalar() or 0
        cancellation_rate = round((cancelled_bookings / total_bookings * 100) if total_bookings else 0, 1)
        avg_rating = db.query(func.avg(TourReview.rating)).filter(TourReview.tour_id == tour.id, TourReview.status == "approved").scalar()

        data.append({
            "tour_name": tour.title,
            "supplier": tour.supplier.supplier_name if tour.supplier else None,
            "destination": ", ".join(filter(None, [tour.city.city_name if tour.city else None, tour.country.country_name if tour.country else None])),
            "category": tour.category.category_name if tour.category else None,
            # No page-view or tour-scoped enquiry tracking exists anywhere in
            # the codebase today - shown honestly as "Not tracked" rather
            # than a fabricated number, same for the conversion rate that
            # depends on views.
            "views": "Not tracked",
            "enquiries": "Not tracked",
            "bookings": total_bookings,
            "confirmed_travellers": confirmed_travellers,
            "booking_conversion_rate": "Not tracked",
            "cancellation_rate": f"{cancellation_rate}%",
            "revenue": _money(revenue),
            "average_rating": round(float(avg_rating), 1) if avg_rating is not None else None,
        })
    return {"status": "success", "data": data}


@router.get("/cancellation-refund-report")
def cancellation_refund_report(
    params: dict = Depends(_period_params),
    supplier_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.admin")),
):
    _require_admin_report(current_user)
    query = db.query(Booking).filter(Booking.booking_status == "cancelled").options(
        joinedload(Booking.customer), joinedload(Booking.supplier)
    )
    query = _apply_range(query, Booking.cancelled_at, params["start"], params["end"])
    if supplier_id:
        query = query.filter(Booking.supplier_id == supplier_id)
    bookings = query.order_by(Booking.cancelled_at.desc()).limit(ROW_REPORT_LIMIT).all()

    booking_ids = [b.id for b in bookings]
    requests_by_booking = {}
    if booking_ids:
        for req in db.query(CancellationRequest).filter(CancellationRequest.booking_id.in_(booking_ids)).all():
            requests_by_booking[req.booking_id] = req

    canceller_ids = [b.cancelled_by for b in bookings if b.cancelled_by]
    canceller_roles: dict[int, str] = {}
    if canceller_ids:
        for user in db.query(User).options(joinedload(User.role)).filter(User.id.in_(canceller_ids)).all():
            canceller_roles[user.id] = user.role.slug if user.role else "-"

    data = []
    for b in bookings:
        req = requests_by_booking.get(b.id)
        refundable_amount = req.refund_amount if req else None
        cancellation_charge = (float(b.final_amount or 0) - float(refundable_amount)) if refundable_amount is not None else None
        cancelled_by_role = canceller_roles.get(b.cancelled_by, "-")
        data.append({
            "booking_id": b.id,
            "tour": b.tour_name,
            "customer": b.customer.full_name if b.customer else None,
            "supplier": b.supplier.supplier_name if b.supplier else b.supplier_name,
            "cancellation_date": b.cancelled_at,
            "cancelled_by": cancelled_by_role,
            "cancellation_reason": b.cancellation_reason or (req.reason if req else None),
            "booking_amount": _money(b.final_amount),
            "cancellation_charge": _money(cancellation_charge) if cancellation_charge is not None else "-",
            "refundable_amount": _money(refundable_amount) if refundable_amount is not None else "-",
            "refund_status": req.status if req else ("refunded" if b.payment_status in {"refunded", "partially_refunded"} else "-"),
            "refund_date": req.refund_processed_at if req else None,
        })
    return {"status": "success", "data": data}


# ---------------------------------------------------------------------------
# Supplier-facing reports - always scoped to the caller's own supplier
# profile via get_actor_supplier, never platform-wide.
# ---------------------------------------------------------------------------

def _supplier_report_booking_query(db: Session, supplier_id: int, *, period: str, start_date: str, end_date: str, booking_status: str = "", payment_status: str = ""):
    """Base query for supplier self-service reports.

    The supplier id is always resolved server-side.  Period filters use the
    booking creation date, making weekly/monthly/quarterly totals consistent
    with the individual booking rows and CSV exports.
    """
    if period not in PERIODS:
        raise HTTPException(status_code=400, detail=f"period must be one of {sorted(PERIODS)}")
    start, end = _period_range(period, start_date, end_date)
    query = _apply_range(
        db.query(Booking).filter(Booking.supplier_id == supplier_id),
        Booking.created_at, start, end,
    )
    if booking_status:
        query = query.filter(Booking.booking_status == booking_status.strip().lower())
    if payment_status:
        query = query.filter(Booking.payment_status == payment_status.strip().lower())
    return query

@router.get("/my-bookings")
def my_bookings_report(
    booking_status: str = Query(default=""),
    payment_status: str = Query(default=""),
    period: str = Query(default="all", description="day|week|month|quarter|half_year|year|custom|all"),
    start_date: str = Query(default=""),
    end_date: str = Query(default=""),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.supplier")),
):
    from app.services.supplier_scope import get_actor_supplier
    supplier = get_actor_supplier(db, current_user)
    query = _supplier_report_booking_query(
        db, supplier.id, period=period, start_date=start_date, end_date=end_date,
        booking_status=booking_status, payment_status=payment_status,
    ).options(joinedload(Booking.customer), joinedload(Booking.agent))
    rows = query.order_by(Booking.created_at.desc()).limit(ROW_REPORT_LIMIT).all()

    data = []
    for b in rows:
        adults, children = _adult_child_counts(b)
        data.append({
            "booking_id": b.id,
            "booking_code": b.booking_code,
            "customer_or_agent": (b.agent.agent_name if b.agent else None) or (b.customer.full_name if b.customer else None) or "-",
            "tour_name": b.tour_name,
            "booking_date": b.created_at,
            "travel_date": b.tour_start_date or b.tour_date,
            "adults": adults,
            "children": children,
            # Supplier-price value of the booking, never the customer total
            # (which includes Tourvaa's markup, tax and fees).
            "total_amount": _money(_supplier_booking_value(b)),
            # Amounts are stored in this booking's own operating currency.
            # The portal must receive the code so its display-currency helper
            # never assumes USD and converts an INR booking a second time.
            "currency": b.currency or "USD",
            "booking_status": b.booking_status,
            "payment_status": b.payment_status,
            "special_requests": b.customer_notes,
        })
    return {"status": "success", "data": data}


@router.get("/my-bookings-summary")
def my_bookings_summary(
    booking_status: str = Query(default=""),
    payment_status: str = Query(default=""),
    period: str = Query(default="all", description="day|week|month|quarter|half_year|year|custom|all"),
    start_date: str = Query(default=""),
    end_date: str = Query(default=""),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.supplier")),
):
    """Supplier-only totals and status breakdown for the selected period."""
    from app.services.supplier_scope import get_actor_supplier
    supplier = get_actor_supplier(db, current_user)
    query = _supplier_report_booking_query(
        db, supplier.id, period=period, start_date=start_date, end_date=end_date,
        booking_status=booking_status, payment_status=payment_status,
    )
    total = query.with_entities(func.count(Booking.id)).scalar() or 0
    confirmed = query.filter(Booking.booking_status.in_(["confirmed", "ready_to_travel", "upcoming"])).with_entities(func.count(Booking.id)).scalar() or 0
    completed = query.filter(Booking.booking_status == "completed").with_entities(func.count(Booking.id)).scalar() or 0
    pending = query.filter(Booking.booking_status.in_(["pending_supplier_acceptance", "pending_payment", "pending_supplier_assignment"])).with_entities(func.count(Booking.id)).scalar() or 0
    cancelled = query.filter(Booking.booking_status.in_(["cancelled", "declined", "refunded"])).with_entities(func.count(Booking.id)).scalar() or 0

    # A supplier's realised booking value is only earned after the trip is
    # completed. Pending, accepted and cancelled records remain visible in
    # the count/status cards above, but must not inflate this financial
    # summary or be reported as completed business.
    completed_query = query.filter(Booking.booking_status == "completed")
    amount_column = func.coalesce(Booking.supplier_gross_amount, Booking.final_amount)
    amounts = [
        {"currency": currency or "USD", "booking_value": _money(amount), "bookings": count}
        for currency, count, amount in completed_query.with_entities(
            Booking.currency, func.count(Booking.id), func.coalesce(func.sum(amount_column), 0),
        ).group_by(Booking.currency).all()
    ]
    status_breakdown = [
        {"status": status or "unknown", "count": count}
        for status, count in query.with_entities(Booking.booking_status, func.count(Booking.id)).group_by(Booking.booking_status).all()
    ]
    return {"status": "success", "data": {
        "total_bookings": total,
        "confirmed_bookings": confirmed,
        "completed_bookings": completed,
        "pending_bookings": pending,
        "cancelled_bookings": cancelled,
        "amounts_by_currency": amounts,
        "status_breakdown": status_breakdown,
    }}


@router.get("/my-earnings")
def my_earnings_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.supplier")),
):
    from app.services.supplier_scope import get_actor_supplier
    from app.services.supplier_ledger import get_supplier_statement
    supplier = get_actor_supplier(db, current_user)
    statement = get_supplier_statement(db, supplier.id)
    # Grouped by currency -- a supplier with bookings in more than one
    # currency must never have those totals blended into one unlabeled
    # number (see services.supplier_ledger.get_supplier_statement).
    return {
        "status": "success",
        "data": {
            "by_currency": [
                {
                    "currency": t["currency"],
                    "total_sales": t["total_gross"],
                    "platform_commission": t["total_commission"],
                    # No adjustment/refund-deduction line-item model exists yet
                    # (see the equivalent note on /reports/supplier-payout-report)
                    # - a refunded pending ledger row is simply excluded from the
                    # totals above rather than itemized, so this is honestly 0.
                    "refund_deductions": _money(0),
                    "adjustments": _money(0),
                    "net_earnings": t["total_net_payable"],
                    "paid_earnings": t["total_paid"],
                    "pending_earnings": t["total_pending"],
                }
                for t in statement["totals_by_currency"]
            ],
        },
    }


@router.get("/my-travellers")
def my_travellers_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.view", "reports.supplier")),
):
    from app.models.bookings import BookingTraveller
    from app.services.supplier_scope import get_actor_supplier
    supplier = get_actor_supplier(db, current_user)
    rows = (
        db.query(BookingTraveller, Booking)
        .join(Booking, BookingTraveller.booking_id == Booking.id)
        .filter(Booking.supplier_id == supplier.id)
        .order_by(Booking.created_at.desc())
        .limit(ROW_REPORT_LIMIT)
        .all()
    )
    data = [
        {
            "booking_id": booking.id,
            "traveller_name": traveller.full_name,
            "contact_information": traveller.email or traveller.phone,
            "travel_date": booking.tour_start_date or booking.tour_date,
            "number_of_travellers": booking.total_travellers,
            "pickup_location": traveller.pickup_location,
            "emergency_contact": traveller.emergency_contact,
            "special_requirements": traveller.special_requirements,
            "booking_status": booking.booking_status,
        }
        for traveller, booking in rows
    ]
    return {"status": "success", "data": data}


# export
REPORT_FETCHERS = {
    "summary": lambda db, params, actor: [summary(params, db, actor)["data"]],
    "bookings": lambda db, params, actor: booking_report(params, db, actor)["data"],
    "payments": lambda db, params, actor: payment_report(params, db, actor)["data"],
    "pending-payments": lambda db, params, actor: pending_payments(db, actor)["data"],
    "overdue-payments": lambda db, params, actor: overdue_payments(db, actor)["data"],
    "country-wise": lambda db, params, actor: country_wise(params, db, actor)["data"],
    "cancellations": lambda db, params, actor: cancellations(params, db, actor)["data"],
    "suppliers": lambda db, params, actor: supplier_report(params, db, actor)["data"],
    "agents": lambda db, params, actor: agent_report(params, db, actor)["data"],
    "customers": lambda db, params, actor: customer_report(params, db, actor)["data"],
    "booking-report": lambda db, params, actor: booking_detail_report(
        params, params.get("booking_status", ""), params.get("payment_status", ""),
        None, params.get("supplier_id") or None, db, actor,
    )["data"],
    "sales-revenue-report": lambda db, params, actor: sales_revenue_report(
        params, params.get("granularity", "day"), db, actor,
    )["data"]["time_series"],
    "payment-report": lambda db, params, actor: payment_detail_report(
        params, params.get("payment_status", ""), params.get("gateway", ""), db, actor,
    )["data"],
    "supplier-report": lambda db, params, actor: supplier_detail_report(params, db, actor)["data"],
    "supplier-payout-report": lambda db, params, actor: supplier_payout_report("", None, db, actor)["data"],
    "tour-performance-report": lambda db, params, actor: tour_performance_report(None, None, None, db, actor)["data"],
    "cancellation-refund-report": lambda db, params, actor: cancellation_refund_report(params, None, db, actor)["data"],
    "my-bookings": lambda db, params, actor: my_bookings_report(
        params.get("booking_status", ""), params.get("payment_status", ""), params.get("period", "all"),
        params.get("start_date", ""), params.get("end_date", ""), db, actor,
    )["data"],
    "my-earnings": lambda db, params, actor: my_earnings_report(db, actor)["data"]["by_currency"],
    "my-travellers": lambda db, params, actor: my_travellers_report(db, actor)["data"],
}

REPORT_LABELS = {
    "summary": "Summary",
    "bookings": "Bookings Summary",
    "payments": "Payments Summary",
    "pending-payments": "Pending Payments",
    "overdue-payments": "Overdue Payments",
    "country-wise": "Country-wise Bookings",
    "cancellations": "Cancellations Summary",
    "suppliers": "Suppliers Summary",
    "agents": "Agent Report",
    "customers": "Customer Report",
    "booking-report": "Booking Report",
    "sales-revenue-report": "Sales and Revenue Report",
    "payment-report": "Payment Report",
    "supplier-report": "Supplier Report",
    "supplier-payout-report": "Supplier Payout Report",
    "tour-performance-report": "Tour Performance Report",
    "cancellation-refund-report": "Cancellation and Refund Report",
    "my-bookings": "My Booking Report",
    "my-earnings": "Earnings Report",
    "my-travellers": "Traveller Report",
}

# The consolidated workbook intentionally contains the reports available on
# the admin Reports screen. Partner-only reports remain scoped to their own
# portals and must not be bundled into an administrator export.
ADMIN_WORKBOOK_REPORTS = (
    "summary", "bookings", "payments", "pending-payments", "overdue-payments",
    "country-wise", "cancellations", "suppliers", "agents", "customers",
    "booking-report", "sales-revenue-report", "payment-report", "supplier-report",
    "supplier-payout-report", "tour-performance-report", "cancellation-refund-report",
)


def _spreadsheet_value(value):
    """Return a stable, cell-safe representation for heterogeneous report rows."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str, ensure_ascii=False)
    return value


def _all_reports_workbook(db: Session, params: dict, actor: User) -> io.BytesIO:
    """Build one worksheet per admin report so no visible report is omitted."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="344054")
    thin = Side(style="thin", color="E4E7EC")

    for report_key in ADMIN_WORKBOOK_REPORTS:
        rows = REPORT_FETCHERS[report_key](db, params, actor)
        columns = list(dict.fromkeys(key for row in rows for key in row.keys()))
        sheet = workbook.create_sheet(title=REPORT_LABELS[report_key][:31])
        sheet.sheet_view.showGridLines = False
        sheet["A1"] = REPORT_LABELS[report_key]
        sheet["A1"].font = Font(size=15, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="175CD3")

        if not columns:
            sheet["A3"] = "No data for the selected period"
            sheet.column_dimensions["A"].width = 36
            continue

        for column_index, column in enumerate(columns, start=1):
            cell = sheet.cell(3, column_index, column.replace("_", " ").title())
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max(len(column) + 4, 14), 32)

        for row_index, row in enumerate(rows, start=4):
            for column_index, column in enumerate(columns, start=1):
                cell = sheet.cell(row_index, column_index, _spreadsheet_value(row.get(column)))
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{max(3, len(rows) + 3)}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("/exports")
def exports(
    report: str = Query(default="summary"),
    format: str = Query(default="csv"),
    period: str = Query(default="all"),
    start_date: str = Query(default=""),
    end_date: str = Query(default=""),
    booking_status: str = Query(default=""),
    payment_status: str = Query(default=""),
    gateway: str = Query(default=""),
    supplier_id: int | None = Query(default=None),
    granularity: str = Query(default="day"),
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_any_permission("reports.export")),
):
    if report == "all":
        _require_admin_report(current_user)
        if format.lower() != "xlsx":
            raise HTTPException(status_code=400, detail="The complete report export is available as XLSX only")
    elif report not in REPORT_FETCHERS:
        raise HTTPException(status_code=400, detail=f"Unknown report '{report}'. Valid options: {sorted(REPORT_FETCHERS)}")
    if report != "all" and format.lower() not in {"csv"}:
        raise HTTPException(status_code=400, detail="Only CSV export is currently supported")

    params = _period_params(period, start_date, end_date)
    # Detail-report filters are retained for an export so the download always
    # matches the table an administrator is viewing.
    params.update({
        "period": period,
        "start_date": start_date,
        "end_date": end_date,
        "booking_status": booking_status,
        "payment_status": payment_status,
        "gateway": gateway,
        "supplier_id": supplier_id,
        "granularity": granularity,
    })
    if report == "all":
        workbook = _all_reports_workbook(db, params, current_user)
        label = f"Complete Reports ({period})"
        log_audit(
            db, actor=current_user, action="export_report", entity_type="report", entity_id=0,
            new_values={"label": label, "format": "XLSX", "report": report, "period": period}, request=request,
        )
        db.commit()
        return StreamingResponse(
            iter([workbook.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="all-reports-{period}.xlsx"'},
        )

    rows = REPORT_FETCHERS[report](db, params, current_user)

    buffer = io.StringIO()
    if rows:
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    else:
        buffer.write("No data for the selected period\n")

    label = f"{REPORT_LABELS.get(report, report)} ({period})"
    log_audit(
        db,
        actor=current_user,
        action="export_report",
        entity_type="report",
        entity_id=0,
        new_values={"label": label, "format": "CSV", "report": report, "period": period},
        request=request,
    )
    db.commit()

    buffer.seek(0)
    filename = f"{report}-{period}.csv"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _serialize_schedule(row: ReportSchedule) -> dict:
    return {
        "id": row.id,
        "report_type": row.report_type,
        "cadence": row.cadence,
        "recipient_emails": row.recipient_emails,
        "is_active": row.is_active,
        "created_by": row.created_by,
        "created_at": row.created_at,
        "last_run_at": row.last_run_at,
    }


class ReportScheduleCreate(BaseModel):
    report_type: str
    cadence: str = "weekly"
    recipient_emails: str


@router.get("/schedule")
def list_report_schedules(db: Session = Depends(get_db), _=Depends(require_any_permission("reports.view"))):
    rows = db.query(ReportSchedule).order_by(ReportSchedule.id.desc()).all()
    return {"status": "success", "data": [_serialize_schedule(row) for row in rows]}


@router.post("/schedule")
def create_report_schedule(data: ReportScheduleCreate, db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.export", "reports.view"))):
    if data.report_type not in REPORT_FETCHERS:
        raise HTTPException(status_code=400, detail=f"Unknown report '{data.report_type}'. Valid options: {sorted(REPORT_FETCHERS)}")
    row = ReportSchedule(
        report_type=data.report_type,
        cadence=data.cadence,
        recipient_emails=data.recipient_emails,
        created_by=current_user.id,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return {
        "status": "success",
        "message": f"Schedule saved. The report will be emailed to the listed recipients {data.cadence}.",
        "data": _serialize_schedule(row),
    }


@router.delete("/schedule/{schedule_id}")
def delete_report_schedule(schedule_id: int, db: Session = Depends(get_db), _=Depends(require_any_permission("reports.export", "reports.view"))):
    row = db.query(ReportSchedule).filter(ReportSchedule.id == schedule_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Schedule not found")
    db.delete(row)
    db.commit()
    return {"status": "success", "message": "Schedule deleted"}


@router.get("/snapshot")
def snapshot(db: Session = Depends(get_db), current_user: User = Depends(require_any_permission("reports.view"))):
    _require_admin_report(current_user)
    now = utcnow()
    curr_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if now.month == 1:
        last_start = curr_start.replace(year=now.year - 1, month=12)
    else:
        last_start = curr_start.replace(month=now.month - 1)

    # Booking Performance
    total_bookings = db.query(func.count(Booking.id)).scalar() or 0
    curr_bookings = db.query(func.count(Booking.id)).filter(Booking.created_at >= curr_start).scalar() or 0
    last_bookings = db.query(func.count(Booking.id)).filter(Booking.created_at >= last_start, Booking.created_at < curr_start).scalar() or 0
    booking_change = round(((curr_bookings - last_bookings) / last_bookings * 100) if last_bookings > 0 else 0, 1)

    # Revenue Summary - net of refunds. captured_amount is left unchanged by
    # a refund (only refunded_amount and payment_status change, see
    # services/payments.py process_refund), so it must be subtracted here or
    # a refunded booking's original charge keeps counting as revenue.
    def _net_revenue(*extra_filters) -> float:
        captured = db.query(func.coalesce(func.sum(Payment.captured_amount), 0)).filter(Payment.payment_status.notin_(["voided", "failed"]), *extra_filters).scalar() or 0
        refunded = db.query(func.coalesce(func.sum(Payment.refunded_amount), 0)).filter(Payment.payment_status.notin_(["voided", "failed"]), *extra_filters).scalar() or 0
        return max(0.0, float(captured) - float(refunded))

    total_revenue = _net_revenue()
    curr_revenue = _net_revenue(Payment.created_at >= curr_start)
    last_revenue = _net_revenue(Payment.created_at >= last_start, Payment.created_at < curr_start)
    revenue_change = round(((curr_revenue - last_revenue) / last_revenue * 100) if last_revenue > 0 else 0, 1)

    # Supplier Approval
    total_suppliers = db.query(func.count(Supplier.id)).scalar() or 0
    pending_suppliers = db.query(func.count(Supplier.id)).filter(Supplier.approval_status == "pending").scalar() or 0

    # Agent Sales
    agent_total = db.query(func.count(Booking.id)).filter(Booking.agent_id.isnot(None)).scalar() or 0
    agent_curr = db.query(func.count(Booking.id)).filter(Booking.agent_id.isnot(None), Booking.created_at >= curr_start).scalar() or 0
    agent_last = db.query(func.count(Booking.id)).filter(Booking.agent_id.isnot(None), Booking.created_at >= last_start, Booking.created_at < curr_start).scalar() or 0
    agent_change = round(((agent_curr - agent_last) / agent_last * 100) if agent_last > 0 else 0, 1)

    # Payment Collection
    total_final = float(db.query(func.coalesce(func.sum(Booking.final_amount), 0)).scalar() or 0)
    total_paid = float(db.query(func.coalesce(func.sum(Booking.amount_paid), 0)).scalar() or 0)
    total_pending_amt = float(db.query(func.coalesce(func.sum(Booking.amount_pending), 0)).scalar() or 0)
    collected_pct = round((total_paid / total_final * 100) if total_final > 0 else 0, 1)
    pending_pct = round(100 - collected_pct, 1)

    # Country-wise
    country_count = db.query(func.count(func.distinct(Booking.country_id))).filter(Booking.country_id.isnot(None)).scalar() or 0

    # Recent exports from audit log
    export_logs = (
        db.query(AuditLog)
        .filter(AuditLog.entity_type == "report", AuditLog.action == "export_report")
        .order_by(AuditLog.id.desc())
        .limit(5)
        .all()
    )
    recent_exports = [
        {
            "id": log.id,
            "label": (log.new_values or {}).get("label", "Report Export"),
            "format": (log.new_values or {}).get("format", "CSV").upper(),
            "exported_at": log.created_at,
        }
        for log in export_logs
    ]

    return {
        "status": "success",
        "data": {
            "booking_performance": {
                "total": total_bookings,
                "current_month": curr_bookings,
                "change_pct": booking_change,
            },
            "revenue_summary": {
                "total": _money(total_revenue),
                "total_raw": float(total_revenue),
                "current_month": _money(curr_revenue),
                "change_pct": revenue_change,
            },
            "supplier_approval": {
                "total": total_suppliers,
                "pending": pending_suppliers,
            },
            "agent_sales": {
                "total": agent_total,
                "current_month": agent_curr,
                "change_pct": agent_change,
            },
            "payment_collection": {
                "collected_pct": collected_pct,
                "pending_pct": pending_pct,
                "total_amount": _money(total_final),
                "collected_amount": _money(total_paid),
                "pending_amount": _money(total_pending_amt),
            },
            "country_wise": {
                "country_count": country_count,
            },
            "meta": {
                "report_types": len(REPORT_FETCHERS),
                "scheduled": _scheduled_report_count(db),
                "total_exports": db.query(func.count(AuditLog.id)).filter(AuditLog.entity_type == "report", AuditLog.action == "export_report").scalar() or 0,
            },
            "recent_exports": recent_exports,
        },
    }
